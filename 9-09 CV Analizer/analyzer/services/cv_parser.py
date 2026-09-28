from io import BytesIO
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader

MAX_CV_BYTES = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {'.pdf', '.docx', '.txt'}


class CVParseError(ValueError):
    pass


def extract_text(uploaded_file):
    if not uploaded_file:
        raise CVParseError('Please upload a CV.')
    if uploaded_file.size > MAX_CV_BYTES:
        raise CVParseError('The CV must be 5 MB or smaller.')

    extension = Path(uploaded_file.name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise CVParseError('Unsupported file type. Upload a PDF, DOCX, or TXT file.')

    raw_bytes = uploaded_file.read()
    try:
        if extension == '.pdf':
            text = '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(raw_bytes)).pages)
        elif extension == '.docx':
            document = Document(BytesIO(raw_bytes))
            text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
            text += '\n' + '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        else:
            text = raw_bytes.decode('utf-8-sig', errors='replace')
    except Exception as exc:
        raise CVParseError('This file could not be read. Please check that it is a valid CV file.') from exc

    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', text)
    text = re.sub(r'\s+', ' ', text).strip()
    if not text:
        raise CVParseError('The uploaded CV does not contain readable text.')
    return text[:50000]
